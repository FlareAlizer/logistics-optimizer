"""Таблицы с результатом и выгрузка в Excel."""

from __future__ import annotations

import math
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .excel import DataError
from .models import AnalysisResult, Group, OrderResult

MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "DD.MM.YYYY"
_MONEY_COLUMNS = {
    "Стоимость",
    "Паллетная стоимость",
    "Стоимость отдельно",
    "По отдельности",
    "Экономия",
    "Значение",
}
_DATE_COLUMNS = {"Дата П.П.С.", "Дата отгрузки", "Дата с", "Дата по"}


def orders_frame(result: AnalysisResult) -> pd.DataFrame:
    """Построчная таблица заказов. Первые колонки идут в том же порядке, что в старой выгрузке."""
    rows = []
    for res in result.results:
        order = res.order
        rows.append(
            {
                "Грузополучатель": order.recipient or "",
                "Дата П.П.С.": order.date,
                "Дата отгрузки": order.shipment_date,
                "Паллеты": order.pallets,
                "Вес, кг": order.weight,
                "Город": order.city,
                "Группа": res.group,
                "Стоимость": None if res.is_error else res.cost,
                "Паллетная стоимость": _finite(res.pallet_cost),
                "Метод доставки": res.method.value,
                "Транспорт": res.transport,
                "Адрес доставки": order.delivery_address,
                "Адрес тр. отдела": order.transport_address or "",
                "Номер заказа": order.order_num,
                "Стоимость отдельно": None if res.is_error else res.separate_cost,
            }
        )
    return pd.DataFrame(rows)


def groups_frame(result: AnalysisResult) -> pd.DataFrame:
    rows = [_group_row(group, result.results) for group in result.groups]
    columns = [
        "Группа",
        "Метод доставки",
        "Города",
        "Получатели",
        "Дата с",
        "Дата по",
        "Заказов",
        "Паллеты",
        "Вес, кг",
        "Транспорт",
        "Стоимость",
        "Наценка",
        "По отдельности",
        "Экономия",
        "Паллетная стоимость",
    ]
    return pd.DataFrame(rows, columns=columns)


def _group_row(group: Group, results: list[OrderResult]) -> dict[str, object]:
    orders = [results[i].order for i in group.indices]
    dates = sorted(order.date for order in orders if order.date is not None)
    return {
        "Группа": group.id,
        "Метод доставки": group.method.value,
        "Города": ", ".join(dict.fromkeys(order.city for order in orders)),
        "Получатели": ", ".join(dict.fromkeys(order.recipient or "" for order in orders)),
        "Дата с": dates[0] if dates else None,
        "Дата по": dates[-1] if dates else None,
        "Заказов": len(orders),
        "Паллеты": sum(order.pallets for order in orders),
        "Вес, кг": sum(order.weight for order in orders),
        "Транспорт": str(group.trucks),
        "Стоимость": group.cost,
        "Наценка": group.markup if group.markup != 1.0 else None,
        "По отдельности": group.separate_cost,
        "Экономия": group.savings,
        "Паллетная стоимость": _finite(group.pallet_cost),
    }


def summary_frame(result: AnalysisResult) -> pd.DataFrame:
    summary = result.summary()
    rows: list[tuple[str, object]] = [
        ("Заказов", summary.orders),
        ("Заказов без расчёта", summary.errors),
        ("Групп", summary.groups),
        ("Прямых рейсов", summary.trips),
        ("Машин 10т", summary.trucks_10t),
        ("Машин 20т", summary.trucks_20t),
        ("Стоимость, если везти каждый заказ отдельно", summary.separate_cost),
        ("Стоимость по плану", summary.plan_cost),
        ("Экономия", summary.savings),
        ("Экономия, %", round(summary.savings_share * 100, 2)),
    ]
    for method, (count, cost) in sorted(summary.by_method.items(), key=lambda item: item[0].value):
        rows.append((f"{method.value or 'БЕЗ МЕТОДА'}: заказов", count))
        rows.append((f"{method.value or 'БЕЗ МЕТОДА'}: стоимость", cost))
    for warning in result.warnings:
        rows.append(("Предупреждение", warning))
    return pd.DataFrame(rows, columns=["Показатель", "Значение"])


def export_xlsx(result: AnalysisResult, path: Path | str) -> Path:
    """Сохраняет три листа: заказы, группы, сводка."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    try:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for name, frame in (
                ("Заказы", orders_frame(result)),
                ("Группы", groups_frame(result)),
                ("Сводка", summary_frame(result)),
            ):
                frame.to_excel(writer, sheet_name=name, index=False)
                _style(writer.sheets[name], list(frame.columns))
    except PermissionError as exc:
        raise DataError(
            f"Не удалось записать {path.name}. Закройте файл в Excel и повторите."
        ) from exc
    except OSError as exc:
        raise DataError(f"Не удалось записать {path.name}: {exc.strerror or exc}.") from exc
    return path


def _style(sheet: Worksheet, columns: list[str]) -> None:
    header_fill = PatternFill("solid", start_color="E8ECF1")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    for position, name in enumerate(columns, start=1):
        letter = get_column_letter(position)
        cells = sheet[letter][1:]
        if name in _MONEY_COLUMNS:
            for cell in cells:
                if isinstance(cell.value, float):
                    cell.number_format = MONEY_FORMAT
        elif name in _DATE_COLUMNS:
            for cell in cells:
                cell.number_format = DATE_FORMAT
        longest = max((len(str(cell.value)) for cell in cells if cell.value is not None), default=0)
        sheet.column_dimensions[letter].width = min(max(len(name), longest) + 3, 60)


def _finite(value: float) -> float | None:
    return value if math.isfinite(value) else None
